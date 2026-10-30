"""Transforma a simulação nas fontes que um analista recebe de verdade: bagunçadas.

- Odontologia e Fisioterapia: exportação mensal do sistema de agendamento (CSV ';', Windows-1252).
  O filtro da exportação usa "data >= último dia do mês anterior", então cada arquivo repete o
  último dia do mês anterior (P01). A partir de maio/2026 o sistema ganhou a coluna CANAL (P02).
- Psicologia e Nutrição: planilhas Excel preenchidas à mão, uma aba por mês, com colunas
  diferentes entre os serviços e os problemas P03 a P13.
- Recepção: planilha da fila de espera (P14 e P15).
- Recepção: exportação mensal do tablet do balcão, com a chegada e a chamada de cada paciente e o
  motivo que ele deu ao desmarcar ou ao justificar uma falta (P16 a P18).
- Cadastros limpos dos sistemas: pacientes (com a isenção da taxa social), alunos, supervisores,
  grade de horários, calendário e a tabela da taxa social por procedimento.

Cada problema plantado é contado no manifesto. Os testes conferem que a padronização
encontra exatamente o que foi plantado e recupera a verdade da simulação.

Uso: python -m clinicas.fontes
"""

from __future__ import annotations

import hashlib
import json
import unicodedata
from datetime import date, datetime, time, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from clinicas.config import (
    CLINICAS,
    FERIADOS,
    FIM,
    INICIO,
    PASTA_FONTES,
    PLANTAO_FERIAS,
    REDE,
    SEMENTE,
    SEMESTRES,
    SERVICOS,
    TAXA_SOCIAL,
    TURNOS,
    UNIDADES,
)
from clinicas.nomes import gerar_nomes_unicos
from clinicas.simulacao import fator_salas, horarios, simular

MESES_ABREV = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]

UNIDADE_SISTEMA = {"REC": "CLIN RECIFE", "OLI": "CLIN OLINDA", "JAB": "CLIN JABOATAO"}
SITUACAO_SISTEMA = {"Realizado": "R", "Falta": "F", "Cancelado pelo paciente": "CP", "Cancelado pela clínica": "CC"}
TIPO_SISTEMA = {"Primeira consulta": "1A CONSULTA", "Retorno": "RETORNO"}

# Jeitos de escrever a mesma situação nas planilhas (P03).
SITUACAO_PLANILHA = {
    "Realizado": ["Atendido", "atendido", "ATENDIDO", "Compareceu", "ok"],
    "Falta": ["Faltou", "falta", "F", "Não compareceu", "faltou "],
    "Cancelado pelo paciente": ["Desmarcou", "Paciente desmarcou", "desmarcado"],
    "Cancelado pela clínica": ["Cancelado (supervisor ausente)", "Clínica cancelou"],
}
# Como a recepção digita cada motivo no tablet (P16: texto livre). A primeira forma é a canônica.
MOTIVO_DIGITADO = {
    "Esqueceu": ["Esqueceu", "esqueceu da consulta", "Paciente esqueceu", "não lembrou"],
    "Trabalho": ["Trabalho", "não foi liberado do trabalho", "trabalho", "Horário de trabalho"],
    "Transporte": ["Transporte", "sem passagem", "ônibus atrasou", "Sem transporte"],
    "Doença": ["Doença", "estava doente", "gripe", "doente"],
    "Chuva": ["Chuva", "choveu muito", "rua alagada", "chuva forte"],
    "Família": ["Família", "problema familiar", "filho doente", "Problema na família"],
    "Compromisso pessoal": ["Compromisso pessoal", "compromisso", "Outro compromisso"],
    "Sentiu melhora": ["Sentiu melhora", "melhorou", "Disse que está melhor"],
}

# Proporção de linhas com cada problema nas planilhas.
TAXA = {
    "data_texto": 0.10,      # P04: data digitada como texto dd/mm/aaaa
    "data_texto_curta": 0.04,  # P04: data como texto d/m/aa
    "hora_texto": 0.25,      # P05: horário como texto ("8h", "08:00")
    "nome_variante": 0.12,   # P06: nome do paciente escrito diferente do cadastro
    "prontuario_vazio": 0.04,  # P07 (Psicologia Recife: 0.09)
    "duplicada": 0.015,      # P08: linha copiada duas vezes
    "aluno_variante": 0.10,  # P11: nome do aluno em maiúsculas ou sem acento
    "ano_errado": 0.003,     # P12: ano digitado errado
    "tipo_vazio": 0.08,      # P13b: coluna "Tipo"/"Primeira vez?" em branco
}


def sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c))


def variante_nome(rng: np.random.Generator, nome: str) -> str:
    escolha = int(rng.integers(4))
    if escolha == 0:
        return nome.upper()
    if escolha == 1:
        return sem_acento(nome)
    if escolha == 2:
        return nome.lower()
    return nome.replace(" ", "  ", 1) + " "


def rotulo_aba(dia: date) -> str:
    return f"{MESES_ABREV[dia.month - 1]} {dia.year}"


def nomes_supervisores(rng: np.random.Generator) -> pd.DataFrame:
    total = sum(c.supervisores for c in CLINICAS)
    nomes = gerar_nomes_unicos(rng, total)
    linhas, i = [], 0
    for c in CLINICAS:
        for k in range(c.supervisores):
            linhas.append({
                "codigo": f"{c.unidade}-{c.servico}-S{k + 1}",
                "nome": ("Profa. " if nomes[i].split()[0][-1] == "a" else "Prof. ") + nomes[i],
                "unidade": c.unidade,
                "servico": c.servico,
            })
            i += 1
    return pd.DataFrame(linhas)


def calendario() -> pd.DataFrame:
    linhas = []
    dia = INICIO
    while dia <= FIM:
        fator = fator_salas(dia)
        if dia.weekday() >= 5:
            situacao = "Fim de semana"
        elif dia in FERIADOS:
            situacao = f"Feriado: {FERIADOS[dia]}"
        elif PLANTAO_FERIAS[0] <= dia <= PLANTAO_FERIAS[1]:
            situacao = "Plantão de férias"
        elif fator == 0:
            situacao = "Recesso acadêmico"
        else:
            situacao = "Aberto"
        linhas.append({"data": dia.isoformat(), "situacao": situacao, "fator_salas": fator})
        dia += timedelta(days=1)
    return pd.DataFrame(linhas)


def grade_horarios() -> pd.DataFrame:
    linhas = []
    for c in CLINICAS:
        for turno, inicio, fim in TURNOS:
            linhas.append({
                "unidade": c.unidade,
                "servico": c.servico,
                "turno": turno,
                "inicio": inicio.strftime("%H:%M"),
                "fim": fim.strftime("%H:%M"),
                "salas": c.salas,
                "duracao_min": SERVICOS[c.servico].duracao_min,
                "horarios_por_sala": sum(1 for t, _ in horarios(c.servico) if t == turno),
            })
    return pd.DataFrame(linhas)


def escrever_sistema(atend: pd.DataFrame, sup: dict[str, str], pasta: Path, contagem: dict) -> None:
    """Exportações mensais do sistema (Odontologia e Fisioterapia)."""
    sis = atend[atend["servico"].map(lambda s: SERVICOS[s].origem == "sistema")].copy()
    pasta.mkdir(parents=True, exist_ok=True)
    mes = date(INICIO.year, INICIO.month, 1)
    sobreposicao = 0
    while mes <= FIM:
        proximo = date(mes.year + (mes.month == 12), mes.month % 12 + 1, 1)
        inicio_filtro = mes - timedelta(days=1)  # o bug: inclui o último dia do mês anterior
        parte = sis[(sis["data"] >= max(inicio_filtro, INICIO)) & (sis["data"] < proximo)]
        sobreposicao += int((parte["data"] < mes).sum())
        saida = pd.DataFrame({
            "COD_AGENDAMENTO": parte["id_agendamento"],
            "DT_AGENDA": parte["data"].map(lambda d: d.strftime("%d/%m/%Y")),
            "HR_AGENDA": parte["hora"].map(lambda h: h.strftime("%H:%M")),
            "UNIDADE": parte["unidade"].map(UNIDADE_SISTEMA),
            "ESPECIALIDADE": parte["servico"].map(lambda s: sem_acento(SERVICOS[s].nome).upper()),
            "NR_PRONTUARIO": parte["prontuario"],
            "NM_PACIENTE": parte["paciente"].map(lambda n: sem_acento(n).upper()),
            "MATRICULA": parte["aluno"],
            "NM_SUPERVISOR": parte["supervisor"].map(sup),
            "TP_ATENDIMENTO": parte["tipo"].map(TIPO_SISTEMA),
            "DS_PROCEDIMENTO": parte["procedimento"].map(lambda p: sem_acento(p).upper()),
            "ST_AGENDA": parte["situacao"].map(SITUACAO_SISTEMA),
            "DURACAO": parte["duracao_min"],
        })
        if mes >= date(2026, 5, 1):  # P02: atualização do sistema acrescentou uma coluna
            canais = ["TELEFONE", "PRESENCIAL", "WHATSAPP"]
            saida["CANAL"] = [canais[i % 3] for i in parte["id_agendamento"]]
        if not saida.empty:
            saida.to_csv(pasta / f"agendamentos_{mes:%Y-%m}.csv", sep=";", index=False, encoding="cp1252",
                         lineterminator="\r\n")  # exportação de sistema Windows
        mes = proximo
    contagem["P01_linhas_repetidas_entre_meses"] = sobreposicao
    contagem["P02_arquivos_com_coluna_CANAL"] = sum(1 for p in pasta.glob("*.csv") if p.stem >= "agendamentos_2026-05")


def escrever_planilhas(atend: pd.DataFrame, alunos: pd.DataFrame, sup: dict[str, str], pasta: Path,
                       rng: np.random.Generator, contagem: dict) -> None:
    pasta.mkdir(parents=True, exist_ok=True)
    nome_aluno = dict(zip(alunos["matricula"], alunos["nome"], strict=True))
    plan = atend[atend["servico"].map(lambda s: SERVICOS[s].origem == "planilha")]
    for (unidade, servico), grupo in plan.groupby(["unidade", "servico"], sort=True):
        wb = Workbook()
        wb.remove(wb.active)
        taxa_vazio = 0.09 if (unidade, servico) == ("REC", "PSI") else TAXA["prontuario_vazio"]
        if servico == "PSI":
            cabecalho = ["Data", "Horário", "Paciente", "Prontuário", "Estagiário(a)", "Supervisor(a)", "Tipo",
                         "Situação", "Observação"]
        else:
            cabecalho = ["Dt atendimento", "Hora", "Nome do paciente", "Nº prontuário", "Aluno(a)",
                         "Professor(a)", "Primeira vez?", "Status"]
        mes = date(INICIO.year, INICIO.month, 1)
        while mes <= FIM:
            proximo = date(mes.year + (mes.month == 12), mes.month % 12 + 1, 1)
            linhas_mes = grupo[(grupo["data"] >= mes) & (grupo["data"] < proximo)]
            ws = wb.create_sheet(rotulo_aba(mes))
            ws.append(cabecalho)
            for celula in ws[1]:
                celula.font = Font(bold=True)
                celula.fill = PatternFill("solid", fgColor="DDEBF7")
            linhas = [montar_linha(r, servico, nome_aluno, sup, rng, contagem, taxa_vazio)
                      for r in linhas_mes.itertuples()]
            plantar_conflitos(linhas, rng, contagem)
            for linha in linhas:
                ws.append(linha)
                if rng.random() < TAXA["duplicada"]:  # P08
                    ws.append(list(linha))
                    contagem["P08_linhas_duplicadas"] += 1
                if rng.random() < 0.004:  # P10: linha em branco no meio
                    ws.append([None] * len(cabecalho))
                    contagem["P10_linhas_em_branco"] += 1
            if linhas:  # P09: linha de total no fim da aba
                ws.append(["TOTAL", None, f"{len(linhas)} atendimentos"])
                contagem["P09_linhas_de_total"] += 1
            ws.column_dimensions["A"].width = 13
            ws.column_dimensions["C"].width = 32
            ws.column_dimensions["E"].width = 30
            ws.column_dimensions["F"].width = 32
            mes = proximo
        nome = f"{SERVICOS[servico].nome} - {UNIDADES[unidade].split()[0]}.xlsx"
        wb.save(pasta / nome)


def montar_linha(r, servico: str, nome_aluno: dict, sup: dict, rng: np.random.Generator, contagem: dict,
                 taxa_vazio: float) -> list:
    # Data (P04 e P12).
    dia: date = r.data
    sorteio = rng.random()
    if sorteio < TAXA["ano_errado"]:
        errado = dia.replace(year=dia.year + 36 if dia.year == 2026 else dia.year - 1)
        valor_data: object = datetime.combine(errado, time())
        contagem["P12_ano_digitado_errado"] += 1
    elif sorteio < TAXA["ano_errado"] + TAXA["data_texto"]:
        valor_data = dia.strftime("%d/%m/%Y")
        contagem["P04_data_como_texto"] += 1
    elif sorteio < TAXA["ano_errado"] + TAXA["data_texto"] + TAXA["data_texto_curta"]:
        valor_data = f"{dia.day}/{dia.month}/{dia.year % 100:02d}"
        contagem["P04_data_como_texto"] += 1
    else:
        valor_data = datetime.combine(dia, time())
    # Horário (P05).
    hora: time = r.hora
    if rng.random() < TAXA["hora_texto"]:
        valor_hora: object = f"{hora.hour}h" if hora.minute == 0 else f"{hora.hour}h{hora.minute:02d}"
        contagem["P05_horario_como_texto"] += 1
    else:
        valor_hora = hora
    # Paciente (P06) e prontuário (P07).
    paciente = r.paciente
    if rng.random() < TAXA["nome_variante"]:
        paciente = variante_nome(rng, paciente)
    prontuario = r.prontuario
    if rng.random() < taxa_vazio:
        prontuario = None
        contagem["P07_prontuario_vazio"] += 1
    elif paciente != r.paciente:
        contagem["P06_nome_paciente_variante"] += 1  # só conta o que muda de fato e tem prontuário
    # Aluno (P11): planilha usa o nome, não a matrícula.
    aluno = nome_aluno[r.aluno]
    if rng.random() < TAXA["aluno_variante"]:
        aluno = sem_acento(aluno).upper() if rng.random() < 0.5 else aluno.upper()
        contagem["P11_nome_aluno_variante"] += 1
    # Situação (P03).
    opcoes = SITUACAO_PLANILHA[r.situacao]
    situacao = opcoes[int(rng.integers(len(opcoes)))]
    if situacao not in ("Atendido", "Faltou", "Desmarcou", "Cancelado (supervisor ausente)"):
        contagem["P03_situacao_escrita_diferente"] += 1
    # Tipo de atendimento (P13b: às vezes em branco).
    primeira = r.tipo == "Primeira consulta"
    vazio = rng.random() < TAXA["tipo_vazio"]
    if vazio:
        contagem["P13b_tipo_em_branco"] += 1
    if servico == "PSI":
        tipo = None if vazio else ("1ª vez" if primeira else "Retorno")
        obs = "Encaminhado pela UBS" if primeira and rng.random() < 0.2 else None
        return [valor_data, valor_hora, paciente, prontuario, aluno, sup[r.supervisor], tipo, situacao, obs]
    tipo = None if vazio else ("Sim" if primeira else "Não")
    return [valor_data, valor_hora, paciente, prontuario, aluno, sup[r.supervisor], tipo, situacao]


def plantar_conflitos(linhas: list[list], rng: np.random.Generator, contagem: dict) -> None:
    """P13: horário digitado errado faz o aluno aparecer com dois pacientes ao mesmo tempo."""
    por_aluno: dict[tuple, list[int]] = {}
    for i, linha in enumerate(linhas):
        if isinstance(linha[0], datetime) and linha[0].year in (2025, 2026) and isinstance(linha[1], time):
            por_aluno.setdefault((linha[0], linha[4]), []).append(i)
    for indices in por_aluno.values():
        if len(indices) >= 2 and rng.random() < 0.012:
            a, b = indices[0], indices[1]
            if linhas[a][1] != linhas[b][1]:
                linhas[b][1] = linhas[a][1]
                contagem["P13_conflito_horario_aluno"] += 1


def escrever_fila(pacientes: pd.DataFrame, pasta: Path, rng: np.random.Generator, contagem: dict) -> None:
    variantes_unidade = {"REC": ["Recife", "RECIFE", "Rec"], "OLI": ["Olinda", "OLINDA"],
                         "JAB": ["Jaboatão", "Jaboatao", "Jaboatão dos Guararapes"]}
    variantes_servico = {"ODO": ["Odontologia", "Odonto"], "FIS": ["Fisioterapia", "Fisio"],
                         "PSI": ["Psicologia", "Psicologia ", "psico"], "NUT": ["Nutrição", "Nutricao"]}
    wb = Workbook()
    ws = wb.active
    ws.title = "Fila"
    ws.append(["Unidade", "Serviço", "Prontuário", "Paciente", "Data de inscrição", "Data da 1ª consulta",
               "Situação"])
    for celula in ws[1]:
        celula.font = Font(bold=True)
    for r in pacientes.sort_values(["data_inscricao", "prontuario"]).itertuples():
        unidade = variantes_unidade[r.unidade][int(rng.integers(len(variantes_unidade[r.unidade])))]
        servico = variantes_servico[r.servico][int(rng.integers(len(variantes_servico[r.servico])))]
        if unidade not in ("Recife", "Olinda", "Jaboatão") or servico not in ("Odontologia", "Fisioterapia",
                                                                               "Psicologia", "Nutrição"):
            contagem["P14_unidade_ou_servico_escrito_diferente"] += 1
        primeira = r.data_primeira if isinstance(r.data_primeira, date) else None
        if r.situacao_fila == "Chamado":
            situacao = "Chamado" if rng.random() < 0.6 else "Atendido"
        elif r.situacao_fila == "Desistiu":
            situacao = "Desistiu" if rng.random() < 0.7 else "Desistência"
        else:
            situacao = "Aguardando"
        valor_primeira: object = datetime.combine(primeira, time()) if primeira else None
        if primeira and rng.random() < 0.004:  # P15: 1ª consulta digitada antes da inscrição
            valor_primeira = datetime.combine(r.data_inscricao - timedelta(days=int(rng.integers(3, 40))), time())
            contagem["P15_primeira_consulta_antes_da_inscricao"] += 1
        ws.append([unidade, servico, r.prontuario, r.nome, datetime.combine(r.data_inscricao, time()),
                   valor_primeira, situacao])
    ws.column_dimensions["D"].width = 32
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 18
    wb.save(pasta / "Fila de espera - Recepção.xlsx")


def escrever_recepcao(atend: pd.DataFrame, pasta: Path, rng: np.random.Generator, contagem: dict) -> None:
    """Exportação mensal do tablet da recepção (todas as clínicas): CHEGADA para quem veio
    (com o horário em que foi chamado) e AVISO para quem desmarcou ou justificou a falta."""
    pasta.mkdir(parents=True, exist_ok=True)
    linhas = []
    for r in atend.itertuples():
        base = {"data": r.data.isoformat(), "unidade": UNIDADES[r.unidade].split()[0],
                "servico": SERVICOS[r.servico].nome, "prontuario": r.prontuario}
        if isinstance(r.chegada, str):
            chamada = r.chamada
            if rng.random() < 0.03:  # P17: o aluno chamou o paciente sem tocar no tablet
                chamada = None
                contagem["P17_chamada_em_branco"] += 1
            linhas.append({**base, "registro": "CHEGADA", "chegada": r.chegada, "chamada": chamada,
                           "motivo_informado": None})
            if rng.random() < 0.01:  # P18: passou o cartão duas vezes
                segunda = (datetime.strptime(r.chegada, "%H:%M") + timedelta(minutes=int(rng.integers(1, 4))))
                linhas.append({**base, "registro": "CHEGADA", "chegada": segunda.strftime("%H:%M"),
                               "chamada": chamada, "motivo_informado": None})
                contagem["P18_chegada_registrada_duas_vezes"] += 1
        elif r.justificou:
            formas = MOTIVO_DIGITADO[r.motivo]
            texto = formas[0] if rng.random() < 0.55 else formas[int(rng.integers(1, len(formas)))]
            if texto != formas[0]:
                contagem["P16_motivo_escrito_livre"] += 1
            linhas.append({**base, "registro": "AVISO", "chegada": None, "chamada": None, "motivo_informado": texto})
    df = pd.DataFrame(linhas)
    meses = df["data"].str[:7]
    for mes, parte in df.groupby(meses, sort=True):
        parte.to_csv(pasta / f"recepcao_{mes}.csv", index=False, encoding="utf-8-sig", lineterminator="\n")


def sha256(caminho: Path) -> str:
    """Hash do conteúdo. Para .xlsx, hash dos valores das células: o arquivo em si muda a cada
    gravação (o openpyxl grava a data e a hora dentro do zip), mas o conteúdo não pode mudar."""
    if caminho.suffix != ".xlsx":
        return hashlib.sha256(caminho.read_bytes()).hexdigest()
    from openpyxl import load_workbook

    h = hashlib.sha256()
    wb = load_workbook(caminho, read_only=True)
    for ws in wb.worksheets:
        h.update(ws.title.encode())
        for linha in ws.iter_rows(values_only=True):
            h.update(repr(linha).encode())
    wb.close()
    return h.hexdigest()


def gerar(destino: Path = PASTA_FONTES, semente: int = SEMENTE) -> dict:
    verdade = simular(semente)
    atend, pacientes, alunos = verdade["atendimentos"], verdade["pacientes"], verdade["alunos"]
    rng = np.random.default_rng(semente + 1)
    sup_df = nomes_supervisores(rng)
    sup = dict(zip(sup_df["codigo"], sup_df["nome"], strict=True))

    # Apaga só os arquivos: no OneDrive, apagar a pasta inteira e recriar logo em seguida
    # costuma falhar com "Acesso negado" (a sincronização segura a pasta por alguns segundos).
    if destino.exists():
        for arquivo in destino.rglob("*"):
            if arquivo.is_file():
                arquivo.unlink()
    (destino / "cadastros").mkdir(parents=True, exist_ok=True)
    contagem: dict[str, int] = {k: 0 for k in [
        "P03_situacao_escrita_diferente", "P04_data_como_texto", "P05_horario_como_texto",
        "P06_nome_paciente_variante", "P07_prontuario_vazio", "P08_linhas_duplicadas", "P09_linhas_de_total",
        "P10_linhas_em_branco", "P11_nome_aluno_variante", "P12_ano_digitado_errado",
        "P13_conflito_horario_aluno", "P13b_tipo_em_branco", "P14_unidade_ou_servico_escrito_diferente",
        "P15_primeira_consulta_antes_da_inscricao", "P16_motivo_escrito_livre", "P17_chamada_em_branco",
        "P18_chegada_registrada_duas_vezes"]}

    escrever_sistema(atend, sup, destino / "sistema", contagem)
    escrever_planilhas(atend, alunos, sup, destino / "planilhas", rng, contagem)
    escrever_fila(pacientes, destino, rng, contagem)
    # Gerador próprio para a recepção: os arquivos das versões anteriores não mudam.
    escrever_recepcao(atend, destino / "recepcao", np.random.default_rng(semente + 8), contagem)

    cad = destino / "cadastros"
    cadastro = pacientes[["prontuario", "nome", "sexo", "data_nascimento", "unidade", "servico"]].copy()
    cadastro["isento_taxa"] = pacientes["isento_taxa"].map({True: "Sim", False: "Não"})
    cadastro.to_csv(cad / "pacientes.csv", index=False, encoding="utf-8", lineterminator="\n")
    pd.DataFrame([{"procedimento": k, "valor": v} for k, v in TAXA_SOCIAL.items()]).to_csv(
        cad / "tabela_taxa_social.csv", index=False, encoding="utf-8", lineterminator="\n")
    alunos.to_csv(cad / "alunos.csv", index=False, encoding="utf-8", lineterminator="\n")
    sup_df.to_csv(cad / "supervisores.csv", index=False, encoding="utf-8", lineterminator="\n")
    grade_horarios().to_csv(cad / "grade_horarios.csv", index=False, encoding="utf-8", lineterminator="\n")
    calendario().to_csv(cad / "calendario_rede.csv", index=False, encoding="utf-8", lineterminator="\n")

    # A verdade fica fora das fontes: só os testes e a conferência usam.
    verdade_dir = destino.parent / "verdade"
    verdade_dir.mkdir(exist_ok=True)
    atend.to_csv(verdade_dir / "atendimentos_verdade.csv", index=False, encoding="utf-8", lineterminator="\n")
    pacientes.to_csv(verdade_dir / "pacientes_verdade.csv", index=False, encoding="utf-8", lineterminator="\n")

    manifesto = {
        "rede": REDE,
        "aviso": "Dados 100% fictícios, gerados por simulação. Não representam nenhuma instituição real.",
        "semente": semente,
        "periodo": [INICIO.isoformat(), FIM.isoformat()],
        "semestres": [s[0] for s in SEMESTRES],
        "agendamentos_reais": len(atend),
        "pacientes": len(pacientes),
        "problemas_plantados": contagem,
        "padroes_plantados": {
            "I1": "Faltas maiores na segunda de manhã e na sexta à tarde; Psicologia é o serviço com mais faltas.",
            "I2": "Olinda/Fisioterapia marca primeira consulta só às segundas e quartas e não consegue contato "
                  "com 40% da fila: agenda ociosa e espera longa ao mesmo tempo.",
            "I3": "Recesso (dez-jan) e plantão de julho: atendimentos caem e a espera de quem entra na fila sobe.",
            "I4": "Cada falta é uma hora de prática a menos para o aluno.",
            "I5": "12% dos pacientes têm propensão 3,2 vezes maior a faltar e concentram boa parte das faltas.",
            "I6": "Quanto mais o paciente esperou na fila, maior a chance de faltar à primeira consulta.",
            "I7": "Recife (Odontologia e Psicologia) com procura acima da capacidade: a fila não zera.",
            "M1": "Chuva vira motivo frequente de abril a julho (inverno chuvoso), mais de manhã.",
            "M2": "Trabalho pesa mais na segunda-feira e de manhã.",
            "M3": "Transporte pesa mais em Jaboatão e no horário das 8h.",
            "M4": "Esquecimento pesa mais na Nutrição (retorno a cada 21 dias) e em quem esperou mais de 45 dias.",
            "E1": "Espera na recepção cresce ao longo do turno (cascata), mais na Odontologia.",
            "E2": "Recife espera mais: uma recepção para quatro clínicas. Segunda de manhã, +5 minutos.",
            "F1": "Taxa social só sobre atendimento realizado; 28% dos pacientes são isentos.",
        },
        "arquivos": {str(p.relative_to(destino)).replace("\\", "/"): sha256(p)
                     for p in sorted(destino.rglob("*")) if p.is_file()},
    }
    (destino / "manifesto.json").write_text(json.dumps(manifesto, ensure_ascii=False, indent=2),
                                            encoding="utf-8", newline="\n")
    return manifesto


if __name__ == "__main__":
    m = gerar()
    print(f"{m['agendamentos_reais']} agendamentos, {m['pacientes']} pacientes")
    for chave, valor in m["problemas_plantados"].items():
        print(f"  {chave}: {valor}")
